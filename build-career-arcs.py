"""Build position career-arc analysis, player grids, profile tables and exports.

Run after collect-nfl-careers.py:  python build-career-arcs.py

Outputs
  local/data/careers/<GROUP>.json        trends, milestones, breakout timing and the player grid per position
  local/data/careers/index.json          position list, metric catalogue, method notes
  local/data/profiles/<id>.json          nflCareer section (season tables) merged into each player profile
  local/nfl-career-seasons-2017-2026.csv every player-season, long format
  local/nfl-career-arcs-2017-2026.xlsx   one wide sheet per position (Year 1..10 blocks) plus trend sheets
"""
from __future__ import annotations

import csv
import json
import statistics
from pathlib import Path

from career_metrics import (GROUP_ORDER, GROUPS, IN_PROGRESS_SEASON, LAST_COMPLETED_SEASON, METRICS, MIN_SAMPLE,
                            PFR_ADVANCED_FROM_2018, ROUND_BUCKETS, TIMING_CLASSES, TIMING_YEARS, LOWER_IS_BETTER,
                            catalog, derive, groups_for_position, metric_value)

ROOT = Path(__file__).resolve().parent
DATA = ROOT / 'local' / 'data'
MAX_CY = LAST_COMPLETED_SEASON - 2017 + 1  # 9 completed career years for the 2017 class
GRID_YEARS = IN_PROGRESS_SEASON - 2017 + 1

WIDE = {
    'QB': ['g', 'gs', 'av', 'pass_yds', 'pass_td', 'int', 'anya', 'rush_yds'],
    'RB': ['g', 'av', 'snap_share', 'carries', 'carry_share', 'rush_yds', 'rec', 'rec_yds', 'total_td'],
    'FB': ['g', 'av', 'snap_share', 'carries', 'rec'],
    'WR': ['g', 'av', 'snap_share', 'targets', 'target_share', 'rec', 'rec_yds', 'rec_td'],
    'TE': ['g', 'av', 'snap_share', 'targets', 'target_share', 'rec', 'rec_yds', 'rec_td'],
    'T': ['g', 'gs', 'av', 'snap_share', 'holds', 'false_starts'], 'G': ['g', 'gs', 'av', 'snap_share', 'holds', 'false_starts'], 'C': ['g', 'gs', 'av', 'snap_share', 'holds', 'false_starts'],
    'DE': ['g', 'gs', 'av', 'snap_share', 'tackles', 'tfl', 'sacks', 'pressures'],
    'DT': ['g', 'gs', 'av', 'snap_share', 'tackles', 'tfl', 'sacks', 'pressures'],
    'LB': ['g', 'gs', 'av', 'snap_share', 'tackles', 'tfl', 'sacks', 'pd'],
    'CB': ['g', 'gs', 'av', 'snap_share', 'tackles', 'def_int', 'pd', 'tgt_allowed', 'rating_allowed'],
    'S': ['g', 'gs', 'av', 'snap_share', 'tackles', 'def_int', 'pd', 'tgt_allowed', 'rating_allowed'],
    'K': ['g', 'av', 'fgm', 'fga', 'fg_pct', 'xp_pct'], 'P': ['g', 'av', 'punts', 'punt_avg', 'net_avg'], 'LS': ['g', 'st_snaps'],
}
COUNT_KEYS = ['cmp', 'att', 'pass_yds', 'pass_td', 'int', 'sacks_taken', 'sack_yds', 'pass_epa', 'carries', 'rush_yds',
              'rush_td', 'targets', 'rec', 'rec_yds', 'rec_td', 'rec_fd', 'yac', 'drops', 'brk_tkl', 'tackles', 'tfl',
              'sacks', 'qb_hits', 'pressures', 'hurries', 'ff', 'pd', 'def_int', 'tgt_allowed', 'cmp_allowed',
              'yds_allowed', 'td_allowed', 'm_tkl', 'tackle_att', 'penalties', 'penalty_yds', 'fgm', 'fga', 'fg50',
              'xpm', 'xpa', 'punts', 'punt_yds', 'punt_net_yds', 'in20', 'kr', 'kr_yds', 'pr', 'pr_yds', 'ppr',
              'g', 'gs', 'snap50', 'side_snaps', 'st_snaps', 'qb_wins', 'qb_losses', 'qb_ties', 'fum_lost', 'scrim_yds', 'total_td',
              'av', 'pro_bowl', 'all_pro', 'all_pro_2nd', 'holds', 'false_starts', 'hands', 'opi', 'dpi', 'def_holds', 'rtp',
              'offsides', 'roughness']
TEXT_COLUMNS = {
    'qb_record': ('Record', 'Regular-season record as the starting QB.'),
    'pfrPos': ('Pos', 'Position listed on the Pro Football Reference team roster that season.'),
    'awards': ('Awards', 'AP MVP, Offensive/Defensive Player of the Year, Offensive/Defensive Rookie of the Year, Comeback Player of the Year, Super Bowl MVP and Walter Payton Man of the Year.'),
}
AWARD_SHORT = {'AP MVP': 'MVP', 'AP Offensive Player of the Year': 'OPOY', 'AP Defensive Player of the Year': 'DPOY',
               'AP Offensive Rookie of the Year': 'OROY', 'AP Defensive Rookie of the Year': 'DROY',
               'AP Comeback Player of the Year': 'Comeback POY', 'Super Bowl MVP': 'Super Bowl MVP',
               'Walter Payton Man of the Year': 'Walter Payton MOY'}


def honor_flag(season):
    if season.get('all_pro'):
        return 'AP'
    if season.get('all_pro_2nd'):
        return 'AP2'
    return 'PB' if season.get('pro_bowl') else ''


def flatten(season, side):
    """Group-specific view of a season: resolve the player's side of the ball into usage metrics."""
    s = {k: v for k, v in season.items() if k != 'usage'}
    usage = season.get('usage') or {}
    main = usage.get(side) or {}
    s['side_snaps'] = main.get('snaps', 0)
    s['snap_share'] = main.get('share', 0)
    if main.get('active') is not None and main.get('games'):
        s['snap_active'] = main['active']
    s['snap50'] = main.get('g50', 0)
    s['st_snaps'] = (usage.get('st') or {}).get('snaps', 0)
    return s


def quantile(values, q):
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * q
    low = int(position)
    high = min(low + 1, len(ordered) - 1)
    return ordered[low] + (ordered[high] - ordered[low]) * (position - low)


def r(value, digits=4):
    return None if value is None else round(value, digits)


def summarize(values):
    if not values:
        return {'n': 0}
    return {'n': len(values), 'mean': r(statistics.fmean(values)), 'median': r(quantile(values, .5)),
            'p25': r(quantile(values, .25)), 'p75': r(quantile(values, .75)), 'p90': r(quantile(values, .9)),
            'max': r(max(values))}


def completed(player, cy):
    return player['year'] + cy - 1 <= LAST_COMPLETED_SEASON


def build_group(key, players):
    spec = GROUPS[key]
    side = spec['side']
    for p in players:
        p['flat'] = {s['cy']: flatten(s, side) for s in p['seasons']}
    metrics = [m for m in dict.fromkeys(spec['metrics'])
               if any(metric_value(m, s) not in (None, 0) for p in players for s in p['flat'].values() if not s.get('inProgress'))]
    out = {'key': key, 'label': spec['label'], 'positions': spec['positions'], 'side': side, 'default': spec['default'],
           'metrics': metrics, 'wide': [m for m in WIDE[key] if m in metrics],
           'milestones': [{'label': m['label'], 'metric': m['metric']} for m in spec['milestones']],
           'buckets': {}}
    for bucket, (bucket_label, test) in ROUND_BUCKETS.items():
        members = [p for p in players if test(p['round'])]
        stats = {}
        change = {}
        for metric in metrics:
            rate = METRICS[metric][3] == 'rate'
            bases = {}
            for basis in (['all'] if rate else ['all', 'played']):
                years = []
                for cy in range(1, MAX_CY + 1):
                    cohort = [p for p in members if completed(p, cy)]
                    values = [v for p in cohort if (v := metric_value(metric, p['flat'][cy], basis)) is not None]
                    played = sum(1 for p in cohort if p['flat'][cy]['g'] > 0)
                    years.append({'cy': cy, 'cohort': len(cohort), 'played': played, **summarize(values)})
                bases[basis] = years
            stats[metric] = bases
            deltas = []
            for cy in range(2, MAX_CY + 1):
                pairs = []
                for p in members:
                    if not completed(p, cy):
                        continue
                    a, b = p['flat'][cy - 1], p['flat'][cy]
                    if a['g'] > 0 and b['g'] > 0:
                        va, vb = metric_value(metric, a, 'played'), metric_value(metric, b, 'played')
                        if va is not None and vb is not None:
                            pairs.append(vb - va)
                lower = metric in LOWER_IS_BETTER
                improved = sum(1 for d in pairs if (d < 0 if lower else d > 0))
                deltas.append({'cy': cy, 'n': len(pairs), 'median': r(quantile(pairs, .5)),
                               'improved': r(improved / len(pairs)) if pairs else None})
            change[metric] = deltas
        milestone_rates = []
        timing = []
        for milestone, raw in zip(out['milestones'], spec['milestones']):
            charted = raw['metric'] in PFR_ADVANCED_FROM_2018
            years = []
            for cy in range(1, MAX_CY + 1):
                cohort = [p for p in members if completed(p, cy) and not (charted and p['flat'][cy]['season'] < 2018)]
                hits = sum(1 for p in cohort if p['flat'][cy]['g'] > 0 and raw['test'](p['flat'][cy]))
                years.append({'cy': cy, 'cohort': len(cohort), 'hits': hits, 'rate': r(hits / len(cohort)) if cohort else None})
            milestone_rates.append({**milestone, 'years': years})
            first_class = max(TIMING_CLASSES[0], 2018) if charted else TIMING_CLASSES[0]
            cohort = [p for p in members if first_class <= p['year'] <= TIMING_CLASSES[1]]
            first = [0] * TIMING_YEARS
            for p in cohort:
                for cy in range(1, TIMING_YEARS + 1):
                    s = p['flat'][cy]
                    if s['g'] > 0 and raw['test'](s):
                        first[cy - 1] += 1
                        break
            timing.append({**milestone, 'classes': [first_class, TIMING_CLASSES[1]], 'cohort': len(cohort),
                           'firstYear': first, 'never': len(cohort) - sum(first)})
        # Players who reached the position's first milestone at least once in Years 1-5
        lead = spec['milestones'][0]
        established = {p['id'] for p in members if p['year'] <= TIMING_CLASSES[1]
                       and any(p['flat'][cy]['g'] > 0 and lead['test'](p['flat'][cy]) for cy in range(1, TIMING_YEARS + 1))}
        peaks = {}
        for metric in metrics:
            if METRICS[metric][3] == 'rate':
                continue
            charted = metric in PFR_ADVANCED_FROM_2018
            first_class = max(TIMING_CLASSES[0], 2018) if charted else TIMING_CLASSES[0]
            lower = metric in LOWER_IS_BETTER
            result = {}
            for scope in ('all', 'established'):
                cohort = [p for p in members if first_class <= p['year'] <= TIMING_CLASSES[1] and (scope == 'all' or p['id'] in established)]
                counts = [0] * TIMING_YEARS
                for p in cohort:
                    values = [(metric_value(metric, p['flat'][cy], 'all'), cy) for cy in range(1, TIMING_YEARS + 1)]
                    values = [(v, cy) for v, cy in values if v is not None and p['flat'][cy]['g'] > 0]
                    if not values or (not lower and max(v for v, _ in values) <= 0):
                        continue
                    best = min(values, key=lambda x: (x[0], x[1])) if lower else max(values, key=lambda x: (x[0], -x[1]))
                    counts[best[1] - 1] += 1
                result[scope] = {'cohort': len(cohort), 'counts': counts, 'none': len(cohort) - sum(counts)}
            peaks[metric] = {'classes': [first_class, TIMING_CLASSES[1]], 'establishedBy': lead['label'], **result}
        out['buckets'][bucket] = {'label': bucket_label, 'players': len(members), 'stats': stats, 'change': change,
                                  'milestones': milestone_rates, 'timing': timing, 'peaks': peaks}
    grid = []
    for p in sorted(players, key=lambda p: (p['year'], p['pick'])):
        rows = []
        for cy, s in sorted(p['flat'].items()):
            rows.append([cy, s['season'], s['g'], 1 if s.get('inProgress') else 0, '/'.join(s.get('teams', [])), honor_flag(s)] +
                        [metric_value(m, s, 'played') if s['g'] > 0 else None for m in metrics])
        grid.append({'id': p['id'], 'name': p['name'], 'year': p['year'], 'round': p['round'], 'pick': p['pick'],
                     'team': p['team'], 'pos': p['position'], 'college': p['college'], 's': rows})
    out['gridColumns'] = ['cy', 'season', 'g', 'inProgress', 'teams', 'honor'] + metrics
    out['players'] = grid
    return out


def career_totals(seasons):
    total = {}
    for s in seasons:
        for key in COUNT_KEYS:
            if isinstance(s.get(key), (int, float)):
                total[key] = total.get(key, 0) + s[key]
    for key in ('scrim_yds', 'total_td', 'ypc', 'ypr', 'catch_pct', 'ypt', 'cmp_pct', 'ypa', 'anya', 'rating'):
        total.pop(key, None)
    derive(total)
    if total.get('tgt_allowed'):
        total.pop('rating_allowed', None)
    if total.get('tackle_att'):
        total['missed_tkl_pct'] = round((total.get('m_tkl') or 0) / total['tackle_att'], 4)
    gs = total.get('gs')
    if gs and 'qb_wins' in total:
        total['win_pct'] = round((total['qb_wins'] + .5 * total.get('qb_ties', 0)) / gs, 4)
    return total


def profile_tables(record, career, groups):
    sides = []
    tables = []
    for key in groups:
        spec = GROUPS[key]
        if spec['side'] not in sides:
            sides.append(spec['side'])
        for label, keys in spec['tables']:
            if label not in [t[0] for t in tables]:
                tables.append((label, keys))
    categories = []
    side_labels = {'off': 'offensive', 'def': 'defensive', 'st': 'special-teams'}
    for side in sides:
        flat = [flatten(s, side) for s in career['seasons']]
        keys = ['pfrPos', 'g', 'gs', 'snap50', 'snap_share', 'snap_active', 'side_snaps'] + (['st_snaps'] if side != 'st' else [])
        if side == 'st':
            keys = ['pfrPos', 'g', 'gs', 'side_snaps']
        if 'QB' in groups:
            keys = keys[:3] + ['qb_record'] + keys[3:]
        categories.append(category('Playing time' + (' (' + side_labels[side] + ')' if len(sides) > 1 else ''), keys, flat, side))
    flat = [flatten(s, sides[0]) for s in career['seasons']]
    for label, keys in tables:
        categories.append(category(label, keys, flat, sides[0]))
    return [c for c in categories if c]


def category(label, keys, flat, side):
    shown = []
    for key in keys:
        if key == 'qb_record':
            if any(s.get('qb_wins') is not None for s in flat):
                shown.append(key)
            continue
        if any(s['g'] > 0 and s.get(key) not in (None, 0) for s in flat):
            shown.append(key)
    if not shown or (label not in ('Playing time',) and not label.startswith('Playing time') and not any(s['g'] for s in flat)):
        return None
    columns = []
    for key in shown:
        if key in TEXT_COLUMNS:
            columns.append({'key': key, 'label': TEXT_COLUMNS[key][0], 'format': 'text', 'description': TEXT_COLUMNS[key][1]})
        else:
            spec = METRICS[key]
            description = spec[5]
            if key in ('snap_share', 'snap_active', 'snap50', 'side_snaps'):
                description += ' Side: ' + {'off': 'offense', 'def': 'defense', 'st': 'special teams'}[side] + '.'
            columns.append({'key': key, 'label': spec[1], 'title': spec[0], 'format': spec[2], 'description': description})
    seasons = []
    for s in flat:
        values = {}
        for key in shown:
            if key == 'qb_record':
                if s.get('qb_wins') is not None:
                    values[key] = f"{s.get('qb_wins', 0)}-{s.get('qb_losses', 0)}" + (f"-{s['qb_ties']}" if s.get('qb_ties') else '')
            elif key == 'awards':
                if s.get('awards'):
                    values[key] = ', '.join(AWARD_SHORT.get(a, a) for a in s['awards'])
            elif key == 'pfrPos':
                if s.get('pfrPos'):
                    values[key] = s['pfrPos']
            elif s['g'] > 0 and s.get(key) is not None:
                values[key] = s[key]
        seasons.append({'season': s['season'], 'cy': s['cy'], 'teams': s.get('teams', []), 'g': s['g'],
                        'inProgress': bool(s.get('inProgress')), 'values': values})
    total = career_totals(flat)
    total_values = {}
    for key in shown:
        if key == 'qb_record':
            if total.get('qb_wins') is not None:
                total_values[key] = f"{total.get('qb_wins', 0)}-{total.get('qb_losses', 0)}" + (f"-{total['qb_ties']}" if total.get('qb_ties') else '')
        elif key == 'awards':
            won = [f"{AWARD_SHORT.get(a, a)} {s['season']}" for s in flat for a in s.get('awards', [])]
            if won:
                total_values[key] = ', '.join(won)
        elif key == 'pfrPos':
            continue
        elif key in total and key not in ('snap_share', 'snap_active'):
            total_values[key] = round(total[key], 4) if isinstance(total[key], float) else total[key]
    return {'label': label, 'columns': columns, 'seasons': seasons, 'career': {'values': total_values}}


def fmt_excel(key, value):
    if value is None:
        return None
    return round(value, 3) if isinstance(value, float) else value


def write_exports(records, careers, group_data):
    long_keys = list(dict.fromkeys(k for key in GROUP_ORDER for k in group_data[key]['metrics']))
    path = ROOT / 'local' / 'nfl-career-seasons-2017-2026.csv'
    with path.open('w', encoding='utf-8', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['Player ID', 'Player', 'Draft position', 'Position group', 'Draft year', 'Round', 'Pick', 'Drafted by',
                         'College', 'Career year', 'Season', 'Season status', 'NFL team(s)'] + [METRICS[k][0] for k in long_keys] + ['PFR player page'])
        for record in records:
            groups = groups_for_position(record['position'], record['role'])
            side = GROUPS[groups[0]]['side']
            career = careers['players'][record['id']]
            for season in career['seasons']:
                s = flatten(season, side)
                status = 'In progress through Week %d' % careers['inProgress']['throughWeek'] if s.get('inProgress') else 'Complete'
                writer.writerow([record['id'], record['name'], record['position'], '/'.join(groups), record['year'], record['round'],
                                 record['pick'], record['team'], record['college'], s['cy'], s['season'], status, '/'.join(s.get('teams', []))] +
                                [fmt_excel(k, metric_value(k, s, 'played')) if s['g'] > 0 else (0 if k == 'g' else None) for k in long_keys] + [career['pfrUrl']])
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print('openpyxl not installed; skipped Excel export')
        return
    wb = Workbook()
    readme = wb.active
    readme.title = 'Read me'
    notes = [
        'NFL Draft Archive: career arcs for every player drafted 2017-2026',
        f'Data as of {careers["asOf"]}. Regular season only. {IN_PROGRESS_SEASON} is in progress through Week {careers["inProgress"]["throughWeek"]}.',
        'Year 1 is the draft season. Each position sheet has one row per drafted player and a block of columns per career year.',
        'Blank cells: season not reached yet, or the player did not play that season (G = 0).',
        'Snap share = share of all team snaps on the player\'s side of the ball across the full season (includes games missed).',
        'Trend sheets: All drafted = every player whose class has completed that career year; seasons without games count as zero for counting stats.',
        'Rate stats (shares, per-attempt figures) only include seasons meeting the qualifier in the metric list below.',
        'Sources: nflverse redistributions of NFL play-by-play statistics, PFR snap counts, PFR advanced charting (2018+) and PFR draft records.',
        '',
        'Metric', 'Definition']
    for line in notes[:-2]:
        readme.append([line])
    readme.append(['Metric', 'Definition'])
    for key, spec in METRICS.items():
        readme.append([spec[0], spec[5]])
    readme['A1'].font = Font(bold=True, size=14)
    readme.column_dimensions['A'].width = 34
    readme.column_dimensions['B'].width = 110
    header_fill = PatternFill('solid', fgColor='152B45')
    year_fill = PatternFill('solid', fgColor='E9EDF3')
    for key in GROUP_ORDER:
        data = group_data[key]
        if not data['players']:
            continue
        wide = data['wide']
        ws = wb.create_sheet(key)
        base = ['Player', 'Pos', 'Draft', 'Rd', 'Pick', 'Drafted by', 'College']
        ws.append(base + [f'Year {cy}' if i == 0 else '' for cy in range(1, GRID_YEARS + 1) for i, _ in enumerate(wide)])
        ws.append([''] * len(base) + [METRICS[m][1] for _ in range(GRID_YEARS) for m in wide])
        for col in range(1, ws.max_column + 1):
            for row in (1, 2):
                cell = ws.cell(row=row, column=col)
                cell.font = Font(bold=True, color='FFFFFF' if row == 1 else '152B45')
                cell.fill = header_fill if row == 1 else year_fill
                cell.alignment = Alignment(horizontal='center')
        for start in range(len(base) + 1, ws.max_column + 1, len(wide)):
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + len(wide) - 1)
        index = {m: data['gridColumns'].index(m) for m in wide}
        for player in data['players']:
            row = [player['name'], player['pos'], player['year'], player['round'], player['pick'], player['team'], player['college']]
            by_cy = {s[0]: s for s in player['s']}
            for cy in range(1, GRID_YEARS + 1):
                s = by_cy.get(cy)
                for m in wide:
                    row.append(fmt_excel(m, s[index[m]]) if s and s[2] > 0 else None)
            ws.append(row)
        for col in range(len(base) + 1, ws.max_column + 1):
            metric = wide[(col - len(base) - 1) % len(wide)]
            ws.column_dimensions[get_column_letter(col)].width = 9
            if METRICS[metric][2] == 'pct':
                for cell in ws.iter_cols(min_col=col, max_col=col, min_row=3):
                    for c in cell:
                        c.number_format = '0.0%'
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['F'].width = 22
        ws.column_dimensions['G'].width = 18
        ws.freeze_panes = 'H3'
        ts = wb.create_sheet(key + ' trends')
        bucket = data['buckets']['all']
        ts.append([f'{data["label"]}: all rounds, completed seasons through {LAST_COMPLETED_SEASON}'])
        ts['A1'].font = Font(bold=True, size=13)
        ts.append(['Metric', 'Statistic'] + [f'Year {cy}' for cy in range(1, MAX_CY + 1)])
        for cell in ts[2]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
        for metric in data['metrics']:
            years = bucket['stats'][metric]['all']
            pct = METRICS[metric][2] == 'pct'
            for label, field in (('Sample (players)', 'n'), ('Average', 'mean'), ('Median', 'median'), ('75th percentile', 'p75'), ('90th percentile', 'p90')):
                row = [METRICS[metric][0] if field == 'n' else '', label] + [y.get(field) for y in years]
                ts.append(row)
                if pct and field != 'n':
                    for c in ts[ts.max_row][2:]:
                        c.number_format = '0.0%'
        ts.append([])
        ts.append(['Milestone', 'Share of drafted players'] + [f'Year {cy}' for cy in range(1, MAX_CY + 1)])
        for cell in ts[ts.max_row]:
            cell.font = Font(bold=True)
        for milestone in bucket['milestones']:
            ts.append([milestone['label'], 'Rate'] + [y['rate'] for y in milestone['years']])
            for c in ts[ts.max_row][2:]:
                c.number_format = '0.0%'
        ts.column_dimensions['A'].width = 34
        ts.column_dimensions['B'].width = 24
        ts.freeze_panes = 'C3'
    wb.save(ROOT / 'local' / 'nfl-career-arcs-2017-2026.xlsx')


def main():
    drafts = json.loads((ROOT / 'drafts.json').read_text(encoding='utf-8'))
    careers = json.loads((DATA / 'nfl-careers.json').read_text(encoding='utf-8'))
    records = [{'id': f'{y}-{row[1]}', 'year': int(y), 'round': row[0], 'pick': row[1], 'team': row[2], 'name': row[4],
                'position': row[5], 'college': row[6]} for y, picks in drafts.items() for row in picks]
    for r in records:
        r['role'] = careers['players'][r['id']].get('nflRole')
    unassigned = [r['position'] for r in records if not groups_for_position(r['position'], r['role'])]
    assert not unassigned, f'Positions without a career group: {sorted(set(unassigned))}'
    out_dir = DATA / 'careers'
    out_dir.mkdir(exist_ok=True)
    group_data = {}
    index = {'asOf': careers['asOf'], 'inProgress': careers['inProgress'], 'lastCompletedSeason': LAST_COMPLETED_SEASON,
             'maxCareerYear': MAX_CY, 'totalPlayers': len(records), 'gridYears': GRID_YEARS, 'minimumSample': MIN_SAMPLE,
             'timing': {'classes': list(TIMING_CLASSES), 'years': TIMING_YEARS},
             'roundBuckets': {k: v[0] for k, v in ROUND_BUCKETS.items()}, 'catalog': catalog(), 'groups': []}
    for key in GROUP_ORDER:
        members = [{**r, 'seasons': careers['players'][r['id']]['seasons']} for r in records if key in groups_for_position(r['position'], r['role'])]
        data = build_group(key, members)
        data.update({'asOf': careers['asOf'], 'inProgress': careers['inProgress']})
        (out_dir / f'{key}.json').write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8', newline='\n')
        group_data[key] = data
        index['groups'].append({'key': key, 'label': data['label'], 'positions': data['positions'], 'players': len(members),
                                'default': data['default']})
        print(f'{key:3} {len(members):4} players, {len(data["metrics"])} metrics, {(out_dir / (key + ".json")).stat().st_size:,} bytes')
    (out_dir / 'index.json').write_text(json.dumps(index, ensure_ascii=False, separators=(',', ':')), encoding='utf-8', newline='\n')
    profiles = {}
    for record in records:
        career = careers['players'][record['id']]
        groups = groups_for_position(record['position'], record['role'])
        played = [s for s in career['seasons'] if s['g'] > 0]
        profiles[record['id']] = {
            'groups': groups, 'pfrUrl': career['pfrUrl'], 'asOf': careers['asOf'], 'inProgress': careers['inProgress'],
            'seasonsPlayed': len([s for s in played if not s.get('inProgress')]), 'games': sum(s['g'] for s in career['seasons']),
            'honors': {'av': sum(s.get('av') or 0 for s in career['seasons']),
                       'avKnown': any('av' in s for s in career['seasons']),
                       'proBowls': [s['season'] for s in career['seasons'] if s.get('pro_bowl')],
                       'allPro': [s['season'] for s in career['seasons'] if s.get('all_pro')],
                       'allPro2': [s['season'] for s in career['seasons'] if s.get('all_pro_2nd')],
                       'proBowlsST': [s['season'] for s in career['seasons'] if s.get('pro_bowl_st')],
                       'allProST': [s['season'] for s in career['seasons'] if s.get('all_pro_st')],
                       'awards': [AWARD_SHORT.get(a, a) + ' ' + str(s['season']) for s in career['seasons'] for a in s.get('awards', [])]},
            'categories': profile_tables(record, career, groups) if played else [],
        }
    written = 0
    for record in records:
        file = DATA / 'profiles' / f'{record["id"]}.json'
        original = file.read_text(encoding='utf-8')
        profile = json.loads(original)
        profile['nflCareer'] = profiles[record['id']]
        published = json.dumps(profile, ensure_ascii=False, separators=(',', ':'))
        if published != original:
            file.write_text(published, encoding='utf-8', newline='\n')
            written += 1
    print(f'Updated the NFL career section in {written} profiles')
    write_exports(records, careers, group_data)
    print('Built career arcs, profile tables, CSV and Excel exports')


if __name__ == '__main__':
    main()
